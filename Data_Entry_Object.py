from openpyxl import Workbook
import docx

# only issue with this class is that it only works with docx and it will break if
# a column has more than one word, have to figure out how to account for that.
class DataEntry:
    def __init__(self, file, column_labels, file_name):
        self.file = file
        self.column_labels = column_labels
        self.file_name = file_name

    # appends the text from doc to list
    def get_doc_text(self):
        doc = docx.Document(self.file)
        full_text = []
        for line in doc.paragraphs:
            full_text.append(line.text)
        return full_text

    # inserts the data from each value of full_text to the next empty row
    # had to intiate create the work book in execute and activate it in this method otherwise it
    # would create a blank XL sheet
    def insert_values(self, doc_text, wb):
        # activate workbook
        ws = wb.active
        # populates first row with the labels
        ws.append(self.column_labels)

        # inserts the rest of the data
        for data in doc_text:
            data = data.split()
            ws.append(data)

    # executes all the above methods
    def execute(self):
        #create a work book
        wb = Workbook()
        task = DataEntry(self.file, self.column_labels, self.file_name)
        full_text = task.get_doc_text()
        task.insert_values(full_text, wb)
        wb.save("{0}.xlsx".format(self.file_name))
